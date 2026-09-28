"""报告生成：依据检测结果产出 Markdown / PDF / Word 三种格式。

数据来源：outputs/results/{result_id}.json（检测器落盘）+ agent.memory 检测档案。
模板：tools/templates/report.md.j2；PDF 用 reportlab（注册 Windows 中文字体）；
Word 用 python-docx。所有产物集中到 outputs/reports/{result_id}/ 下。
"""
import json
import logging
import shutil
from datetime import datetime
from pathlib import Path

from agent import memory
from agent.config import CFG
from tools.sar_knowledge import interpretation_hints

log = logging.getLogger(__name__)

_FONT_CANDIDATES = [
    ("SimHei", "C:/Windows/Fonts/simhei.ttf", None),
    ("MSYaHei", "C:/Windows/Fonts/msyh.ttc", 0),
    ("DengXian", "C:/Windows/Fonts/DENG.TTF", None),
]


def _load_result(result_id: str) -> dict:
    rec = memory.get_detection(result_id)
    if rec is None:
        raise KeyError(f"检测结果不存在：{result_id}（请先执行检测）")
    result_json = CFG.results_dir / f"{result_id}.json"
    data = json.loads(result_json.read_text(encoding="utf-8"))
    data["created_at"] = rec.get("created_at")
    return data


def _render_context(data: dict, interpretation: str | None) -> dict:
    summary = data.get("summary", {})
    if not interpretation:
        interpretation = interpretation_hints(summary.get("num_objects", 0),
                                              summary.get("conf_mean", 0.0))
    params = data.get("params", {})
    cls_counts = summary.get("cls_counts", {})
    return {
        "result_id": data["result_id"],
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "image_path": data.get("image_path", ""),
        "weights": data.get("weights", ""),
        "weights_source": data.get("weights_source", ""),
        "params": params,
        "summary": summary,
        "cls_counts_str": "、".join(f"{k}×{v}" for k, v in cls_counts.items()) or "无",
        "boxes": data.get("boxes", []),
        "interpretation": [p.strip() for p in interpretation.split("\n") if p.strip()],
        "hist_name": Path(data["charts"]["hist"]).name if data.get("charts") else "",
        "position_name": Path(data["charts"]["position"]).name if data.get("charts") else "",
    }


def _collect_assets(result_id: str, data: dict, out_dir: Path) -> dict:
    """把标注图与统计图复制进报告目录，返回供各格式引用的本地文件名。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    names = {}
    for key, src in [("annotated", data.get("vis_path")),
                     ("hist", (data.get("charts") or {}).get("hist")),
                     ("position", (data.get("charts") or {}).get("position"))]:
        if src and Path(src).exists():
            dst = out_dir / f"{result_id}_{key}{Path(src).suffix}"
            shutil.copyfile(src, dst)
            names[key] = dst.name
    return names


def _build_markdown(ctx: dict, assets: dict, out_dir: Path) -> Path:
    from jinja2 import Environment, FileSystemLoader
    env = Environment(loader=FileSystemLoader(str(Path(__file__).parent / "templates")),
                      keep_trailing_newline=True)
    md = env.get_template("report.md.j2").render(
        **ctx,
        annotated_name=assets.get("annotated", "(标注图缺失)"),
    )
    path = out_dir / f"{ctx['result_id']}_report.md"
    path.write_text(md, encoding="utf-8")
    return path


def _register_cn_font():
    """注册中文字体，返回字体名；Windows 自带字体缺失时回退 Helvetica。"""
    from reportlab.pdfbase import pdfmetrics, ttfonts
    for name, path, sub in _FONT_CANDIDATES:
        if not Path(path).exists():
            continue
        try:
            font = (ttfonts.TTFont(name, path, subfontIndex=sub)
                    if sub is not None else ttfonts.TTFont(name, path))
            pdfmetrics.registerFont(font)
            return name
        except Exception as e:  # 字体损坏时尝试下一个
            log.warning("注册字体 %s 失败：%s", path, e)
    return "Helvetica"


def _build_pdf(ctx: dict, assets: dict, out_dir: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_cn_font()
    h1 = ParagraphStyle("h1", fontName=font, fontSize=16, leading=22, spaceAfter=8)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=13, leading=18,
                        spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=16,
                          wordWrap="CJK")

    def img(name, width=12 * cm):
        p = out_dir / name
        if not p.exists():
            return None
        return RLImage(str(p), width=width, height=width * 0.72)

    story = [Paragraph("SAR 图像舰船检测分析报告", h1),
             Paragraph(f"报告编号：{ctx['result_id']}<br/>生成时间：{ctx['generated_at']}"
                       f"<br/>图像文件：{ctx['image_path']}<br/>"
                       f"检测模型：{ctx['weights']}（{ctx['weights_source']}）<br/>"
                       f"检测参数：conf={ctx['params'].get('conf')} / "
                       f"iou={ctx['params'].get('iou')} / "
                       f"imgsz={ctx['params'].get('imgsz')} / "
                       f"device={ctx['params'].get('device')}", body),
             Paragraph("一、检测统计", h2)]
    stat_rows = [["指标", "数值"],
                 ["检出目标数", str(ctx["summary"].get("num_objects", 0))],
                 ["类别计数", ctx["cls_counts_str"]],
                 ["平均/最高/最低置信度",
                  f"{ctx['summary'].get('conf_mean')} / {ctx['summary'].get('conf_max')}"
                  f" / {ctx['summary'].get('conf_min')}"],
                 ["推理耗时", f"{ctx['summary'].get('elapsed_sec')} 秒"]]
    tbl = Table(stat_rows, colWidths=[6 * cm, 8 * cm])
    tbl.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font),
                             ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                             ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story += [tbl, Spacer(1, 8)]
    for key in ("annotated", "hist", "position"):
        im = img(assets.get(key, "_none_"), 12 * cm if key == "annotated" else 9 * cm)
        if im is not None:
            story += [im, Spacer(1, 6)]
    story.append(Paragraph("二、目标明细（按置信度降序）", h2))
    detail = [["#", "类别", "置信度", "bbox (x1,y1,x2,y2)"]]
    detail += [[str(i + 1), b["cls_name"], str(b["conf"]),
                ", ".join(str(v) for v in b["xyxy"])]
               for i, b in enumerate(ctx["boxes"][:30])] or \
              [["-", "-", "-", "（无检出目标）"]]
    dtbl = Table(detail, colWidths=[1 * cm, 2.5 * cm, 2 * cm, 8.5 * cm])
    dtbl.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font),
                              ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                              ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                              ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story += [dtbl, Paragraph("三、智能解读", h2)]
    story += [Paragraph(p, body) for p in ctx["interpretation"]]
    story.append(Paragraph("四、方法说明", h2))
    story.append(Paragraph(
        "数据集：HRSID_OPENSSDD v2（HRSID 与 SSDD/OpenSSDD 合并舰船数据集，640×640，"
        "单一类别 ship）。检测器：YOLOv8（Ultralytics）。统计结论来自检测器结构化输出；"
        "文字解读由 GLM 大模型基于该结构化结果生成，不直接观测原图，保证数字可溯源。"
        "SAR 小目标存在漏检可能，低置信度框存在虚警可能；本报告仅用于学术与民用场景分析。",
        body))
    path = out_dir / f"{ctx['result_id']}_report.pdf"
    SimpleDocTemplate(str(path), pagesize=A4,
                      leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=1.8 * cm, bottomMargin=1.8 * cm).build(story)
    return path


def _build_docx(ctx: dict, assets: dict, out_dir: Path) -> Path:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.add_heading("SAR 图像舰船检测分析报告", level=0)
    meta = doc.add_paragraph()
    for line in (f"报告编号：{ctx['result_id']}",
                 f"生成时间：{ctx['generated_at']}",
                 f"图像文件：{ctx['image_path']}",
                 f"检测模型：{ctx['weights']}（{ctx['weights_source']}）",
                 f"检测参数：conf={ctx['params'].get('conf')} / iou={ctx['params'].get('iou')}"
                 f" / imgsz={ctx['params'].get('imgsz')} / device={ctx['params'].get('device')}"):
        meta.add_run(line + "\n")

    doc.add_heading("一、检测统计", level=1)
    rows = [("检出目标数", str(ctx["summary"].get("num_objects", 0))),
            ("类别计数", ctx["cls_counts_str"]),
            ("平均/最高/最低置信度",
             f"{ctx['summary'].get('conf_mean')} / {ctx['summary'].get('conf_max')}"
             f" / {ctx['summary'].get('conf_min')}"),
            ("推理耗时", f"{ctx['summary'].get('elapsed_sec')} 秒")]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "指标", "数值"
    for k, v in rows:
        c = table.add_row().cells
        c[0].text, c[1].text = k, v

    for title, key, width in (("标注结果", "annotated", 15),
                              ("置信度分布", "hist", 11),
                              ("目标位置分布", "position", 11)):
        if key in assets:
            doc.add_heading(title, level=2)
            doc.add_picture(str(out_dir / assets[key]), width=Cm(width))
            doc.paragraphs[-1].alignment = 1

    doc.add_heading("二、目标明细（按置信度降序）", level=1)
    dt = doc.add_table(rows=1, cols=4)
    dt.style = "Table Grid"
    for i, h in enumerate(("#", "类别", "置信度", "bbox (x1,y1,x2,y2)")):
        dt.rows[0].cells[i].text = h
    if not ctx["boxes"]:
        c = dt.add_row().cells
        c[0].text, c[3].text = "-", "（无检出目标）"
    for i, b in enumerate(ctx["boxes"][:30]):
        c = dt.add_row().cells
        c[0].text = str(i + 1)
        c[1].text = str(b["cls_name"])
        c[2].text = str(b["conf"])
        c[3].text = ", ".join(str(v) for v in b["xyxy"])

    doc.add_heading("三、智能解读", level=1)
    for p in ctx["interpretation"]:
        doc.add_paragraph(p)
    doc.add_heading("四、方法说明", level=1)
    doc.add_paragraph(
        "数据集：HRSID_OPENSSDD v2（HRSID 与 SSDD/OpenSSDD 合并舰船数据集，640×640，"
        "单一类别 ship）。检测器：YOLOv8（Ultralytics）。统计结论来自检测器结构化输出；"
        "文字解读由 GLM 大模型基于该结构化结果生成，不直接观测原图，保证数字可溯源。"
        "SAR 小目标存在漏检可能，低置信度框存在虚警可能；本报告仅用于学术与民用场景分析。")
    path = out_dir / f"{ctx['result_id']}_report.docx"
    doc.save(str(path))
    return path


def generate(result_id: str, fmt: str = "markdown",
             interpretation: str | None = None) -> dict:
    """生成报告；返回 {"ok", "path", "format"} 或 {"ok": False, "error"}。"""
    try:
        fmt = (fmt or "markdown").lower()
        aliases = {"md": "markdown", "word": "docx"}
        fmt = aliases.get(fmt, fmt)
        if fmt not in CFG.report_formats:
            return {"ok": False, "error": f"不支持的报告格式：{fmt}（支持 {CFG.report_formats}）"}
        data = _load_result(result_id)
        out_dir = CFG.reports_dir / result_id
        assets = _collect_assets(result_id, data, out_dir)
        ctx = _render_context(data, interpretation)
        builders = {"markdown": _build_markdown, "pdf": _build_pdf, "docx": _build_docx}
        path = builders[fmt](ctx, assets, out_dir)
        memory.save_report_paths(result_id, fmt, str(path))
        return {"ok": True, "format": fmt, "path": str(path)}
    except KeyError as e:
        return {"ok": False, "error": str(e)}
    except Exception as e:
        log.exception("报告生成失败")
        return {"ok": False, "error": f"报告生成失败：{e}"}
